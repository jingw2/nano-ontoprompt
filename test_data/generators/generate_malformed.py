"""Generates malformed-file fixtures: corrupted OOXML zips and a truncated PDF.

Corruption is done by truncating a valid file's last 300 bytes, which removes
the ZIP end-of-central-directory record (docx/xlsx) without leaving the file
empty — this is what makes MarkItDown's ZipConverter report a genuine
"[ERROR] Invalid or corrupted zip file" instead of a generic not-found error.

python-docx/openpyxl both stamp docProps/core.xml's created/modified fields
with datetime.now(), and python's zipfile module stamps each ZIP entry's own
local-file-header timestamp from wall-clock time at save. Both sources of
wall-clock time are neutralized below: fixed core properties, a regex patch
of docProps/core.xml's <dcterms:modified> (openpyxl's save_workbook() at
writer/excel.py re-stamps `modified` with datetime.now() unconditionally at
save time, clobbering any pre-save assignment — plain property assignment
alone is not sufficient), and a full zip rewrite with a fixed
ZipInfo.date_time.

openpyxl also internally dedupes styles via hash-based containers, so the
serialized order of xl/styles.xml depends on Python's per-process string
hash randomization (PYTHONHASHSEED). The re-exec below pins that seed before
any of openpyxl's code runs, so styles.xml comes out byte-identical across
separate invocations too.
"""
import os
import re
import sys

if os.environ.get("PYTHONHASHSEED") != "0":
    os.environ["PYTHONHASHSEED"] = "0"
    os.execv(sys.executable, [sys.executable] + sys.argv)

import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
import openpyxl
from pypdf import PdfWriter

BASE = Path(__file__).resolve().parent.parent / "edge_cases" / "malformed"
BASE.mkdir(parents=True, exist_ok=True)

FIXED_DATETIME = datetime(2026, 1, 1, 0, 0, 0)
FIXED_ZIP_DATE_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_ISO_DATETIME = b"2026-01-01T00:00:00Z"
_MODIFIED_RE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normalize_zip_timestamps(path: Path) -> None:
    """Rewrites every ZIP entry's local-file-header timestamp to a fixed
    value, and patches docProps/core.xml's <dcterms:modified> content (which
    openpyxl re-stamps with wall-clock time at save time regardless of what
    was assigned beforehand), so byte-for-byte output no longer depends on
    wall-clock time."""
    with zipfile.ZipFile(path, "r") as zf:
        entries = [(info, zf.read(info.filename)) for info in zf.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in entries:
            info.date_time = FIXED_ZIP_DATE_TIME
            if info.filename == "docProps/core.xml":
                data = _MODIFIED_RE.sub(
                    lambda m: m.group(1) + FIXED_ISO_DATETIME + m.group(2), data
                )
            zf.writestr(info, data)


def _truncate(src: Path, dst: Path, cut_bytes: int = 300) -> None:
    data = src.read_bytes()
    dst.write_bytes(data[:-cut_bytes])
    src.unlink()


def generate_corrupted_docx() -> None:
    valid = BASE / "_valid_docx_scratch.docx"
    doc = Document()
    doc.add_paragraph("测试内容 for corruption fixture")
    doc.core_properties.created = FIXED_DATETIME
    doc.core_properties.modified = FIXED_DATETIME
    doc.save(valid)
    _normalize_zip_timestamps(valid)
    _truncate(valid, BASE / "corrupted_docx.docx")


def generate_corrupted_xlsx() -> None:
    valid = BASE / "_valid_xlsx_scratch.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["客户ID", "金额"])
    ws.append(["C001", 100])
    wb.properties.created = FIXED_DATETIME
    wb.properties.modified = FIXED_DATETIME
    wb.save(valid)
    _normalize_zip_timestamps(valid)
    _truncate(valid, BASE / "corrupted_xlsx.xlsx")


def generate_truncated_pdf() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    valid = BASE / "_valid_pdf_scratch.pdf"
    with open(valid, "wb") as f:
        writer.write(f)
    data = valid.read_bytes()
    (BASE / "truncated.pdf").write_bytes(data[: len(data) // 2])
    valid.unlink()


def main() -> None:
    generate_corrupted_docx()
    generate_corrupted_xlsx()
    generate_truncated_pdf()
    print(f"Wrote malformed fixtures to {BASE}")


if __name__ == "__main__":
    main()
