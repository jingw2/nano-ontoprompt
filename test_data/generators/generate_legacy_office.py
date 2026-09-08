"""Generates legacy .doc and .ppt missing-format fixtures for every test_data
domain. python-docx/python-pptx can only write modern OOXML (.docx/.pptx);
producing a genuine legacy binary requires converting through headless
LibreOffice, so this script builds a .docx/.pptx first, then shells out to
`soffice --convert-to doc` / `--convert-to ppt`. This conversion step is a
one-time local dependency for regenerating fixtures — the committed .doc/.ppt
output has no runtime dependency on LibreOffice."""
import shutil
import subprocess
from pathlib import Path

from docx import Document
from pptx import Presentation

from _domain_schemas import DOMAIN_SCHEMAS

TEST_DATA = Path(__file__).resolve().parent.parent
SOFFICE = shutil.which("soffice")


def _table_text(headers: list[str], rows: list[list]) -> list[str]:
    lines = [" | ".join(headers)]
    for row in rows:
        lines.append(" | ".join(str(v) for v in row))
    return lines


def write_docx_then_doc(domain: str, headers: list[str], rows: list[list]) -> None:
    domain_dir = TEST_DATA / domain
    scratch = domain_dir / f"_{domain}_scratch.docx"
    doc = Document()
    doc.add_heading(f"{domain} 示例数据", level=1)
    for line in _table_text(headers, rows):
        doc.add_paragraph(line)
    doc.save(scratch)
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "doc", "--outdir", str(domain_dir), str(scratch)],
        check=True, capture_output=True,
    )
    scratch.unlink()
    (domain_dir / f"_{domain}_scratch.doc").rename(domain_dir / f"{domain}_sample.doc")


def write_pptx_then_ppt(domain: str, headers: list[str], rows: list[list]) -> None:
    domain_dir = TEST_DATA / domain
    scratch = domain_dir / f"_{domain}_scratch.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"{domain} 示例数据"
    body = slide.placeholders[1].text_frame
    body.text = _table_text(headers, rows)[0]
    for line in _table_text(headers, rows)[1:]:
        body.add_paragraph().text = line
    prs.save(scratch)
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "ppt", "--outdir", str(domain_dir), str(scratch)],
        check=True, capture_output=True,
    )
    scratch.unlink()
    (domain_dir / f"_{domain}_scratch.ppt").rename(domain_dir / f"{domain}_sample.ppt")


def main() -> None:
    if not SOFFICE:
        raise SystemExit(
            "soffice not found on PATH. Install LibreOffice "
            "(e.g. `brew install --cask libreoffice`) to regenerate .doc/.ppt fixtures."
        )
    for domain, (headers, rows) in DOMAIN_SCHEMAS.items():
        write_docx_then_doc(domain, headers, rows)
        write_pptx_then_ppt(domain, headers, rows)
    print(f"Wrote .doc/.ppt fixtures for {len(DOMAIN_SCHEMAS)} domains")


if __name__ == "__main__":
    main()
