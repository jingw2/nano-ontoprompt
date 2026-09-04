"""Closed-schema artifact allowlist and forbidden-value safety scanner.

``ArtifactAllowlist`` is a closed schema: unknown fields are rejected, so no
one can widen it into a raw-content carrier by accident. Forbidden values are
always derived from a loaded ``JourneyManifest`` (never a hand-maintained
literal list); pattern-based checks (email/JWT/bearer/sensitive header keys)
are the only hardcoded detectors, because those are structural, not
fixture-specific.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Literal, Mapping, Sequence

from pydantic import BaseModel, ConfigDict, ValidationError

from .contracts import (
    ArtifactSafetyError,
    ScanResult,
    SanitizedBrowserArtifactRef,
    sha256_text,
)

# `JourneyManifest` is Task 1's real typed manifest, not a second,
# hand-invented type here. Follow the same cross-directory import
# convention Task 1's own `backend/tests/runtime/run_registered_cases.py`
# uses: put `test_data/runtime` on sys.path and import the module directly
# (it is not a package under `backend/`, so `from .contracts import ...`
# cannot reach it).
REPO_ROOT = Path(__file__).resolve().parents[3]
_RUNTIME_DATA_DIR = REPO_ROOT / "test_data" / "runtime"
if str(_RUNTIME_DATA_DIR) not in sys.path:
    sys.path.insert(0, str(_RUNTIME_DATA_DIR))

from journey_registry import JOURNEY_IDS, JourneyManifest, load_journey_manifest  # noqa: E402

SCANNER_VERSION = "1"

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_JWT_RE = re.compile(r"\b[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\b")
_BEARER_RE = re.compile(r"\bBearer\s+\S+", re.IGNORECASE)
_SENSITIVE_HEADER_KEYS = {
    "authorization",
    "x-api-key",
    "api-key",
    "cookie",
    "set-cookie",
    "proxy-authorization",
}

# ---------------------------------------------------------------------------
# Closed schemas.
# ---------------------------------------------------------------------------


class ArtifactAllowlist(BaseModel):
    """The only fields ever allowed into uploaded CI evidence."""

    model_config = ConfigDict(extra="forbid", protected_namespaces=())

    schema_version: int
    run_id: str
    journey_id: str
    fixture_manifest_sha256: str
    model_requested: str
    model_observed: str
    model_origin: str
    model_caller: str
    model_config_version_id: str
    preflight_model_id: str
    logical_model_calls: int
    http_attempts: int
    retry_count: int
    call_timestamps: tuple[str, ...]
    call_kinds: tuple[str, ...]
    semantic_outcomes: tuple[str, ...]
    state_transitions: tuple[str, ...]
    citation_ids: tuple[str, ...]
    tool_trace_ids: tuple[str, ...]
    audit_event_ids: tuple[str, ...]
    backend_trace_ids: tuple[str, ...]
    browser_artifacts: tuple[SanitizedBrowserArtifactRef, ...] = ()


class ScannerFailureSummary(BaseModel):
    """Fixed, closed schema for the on-failure evidence upload."""

    model_config = ConfigDict(extra="forbid")

    schema_version: int
    run_id: str
    # Task 5's own CI contract test constructs this class directly with
    # `status="scanner_failed"` (distinguishing "the scanner itself refused
    # to pass this run" from `ScanResult.status`'s generic "failed", which
    # is a different, non-uploaded type). Renamed from the original
    # `Literal["failed"]` for that reason; nothing else keys off this exact
    # string (`test_artifacts.py` only checks the on-disk key set, never the
    # `status` value).
    status: Literal["scanner_failed"]
    reason_codes: tuple[str, ...]
    category_counts: Mapping[str, int]
    scanner_version: str

    def to_json(self) -> Mapping[str, object]:
        return self.model_dump(mode="json")


# ---------------------------------------------------------------------------
# Forbidden-value derivation.
# ---------------------------------------------------------------------------


_MIN_FORBIDDEN_CELL_LENGTH = 4


def _add_cell_value(values: set[str], text: str) -> None:
    """Add one extracted cell/paragraph value, skipping anything too short to
    be a meaningful forbidden value. A bare "S" or "256" is common enough
    (spreadsheet codes, row counts) that treating it as forbidden would flag
    unrelated, legitimate substrings elsewhere (e.g. inside a field name like
    "fixture_manifest_sha256") rather than an actual echoed fixture cell."""
    text = text.strip()
    if len(text) >= _MIN_FORBIDDEN_CELL_LENGTH:
        values.add(text)


def _extract_csv_cells(path: Path) -> set[str]:
    values: set[str] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            for cell in row:
                _add_cell_value(values, cell)
    return values


def _extract_xlsx_cells(path: Path) -> set[str]:
    from openpyxl import load_workbook

    values: set[str] = set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    _add_cell_value(values, str(cell))
    finally:
        workbook.close()
    return values


def _extract_docx_text(path: Path) -> set[str]:
    from docx import Document

    values: set[str] = set()
    document = Document(str(path))
    for paragraph in document.paragraphs:
        _add_cell_value(values, paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                _add_cell_value(values, cell.text)
    return values


def _extract_input_values(entry: Mapping[str, object]) -> set[str]:
    """Pull the real synthetic cell/paragraph content out of one fixture
    input's referenced file (a real, existing, read-only domain source file
    under ``test_data/...``, per ``entry["path"]``). This is the content a
    real DeepSeek response or browser artifact could actually echo back;
    images are skipped (no OCR here)."""
    rel_path = str(entry.get("path", ""))
    if not rel_path:
        return set()
    full_path = REPO_ROOT / rel_path
    if not full_path.exists():
        return set()
    media_type = str(entry.get("media_type", ""))
    try:
        if media_type == "text/csv":
            return _extract_csv_cells(full_path)
        if media_type.endswith("spreadsheetml.sheet"):
            return _extract_xlsx_cells(full_path)
        if media_type.endswith("wordprocessingml.document"):
            return _extract_docx_text(full_path)
    except Exception:
        # A source file this scanner cannot parse contributes no forbidden
        # values rather than failing manifest loading.
        return set()
    return set()


def _collect_dialogue_values(scenario: Mapping[str, object]) -> set[str]:
    """Every literal string in one dialogue scenario -- the question,
    expected keywords, dialogue id, and outcome labels -- except
    ``citation_source_ids``. Citation IDs are excluded on purpose: they are
    exactly the identifiers ``ArtifactAllowlist.citation_ids`` is meant to
    legitimately disclose (the semantic validator requires a real response to
    cite them), so forbidding them would make every passing scan impossible."""
    values: set[str] = set()

    def walk(node: object, *, skip: bool) -> None:
        if isinstance(node, str):
            if node and not skip:
                values.add(node)
        elif isinstance(node, Mapping):
            for key, value in node.items():
                walk(value, skip=(key == "citation_source_ids"))
        elif isinstance(node, (list, tuple)):
            for item in node:
                walk(item, skip=skip)

    walk(scenario, skip=False)
    return values


def _collect_pii_secret_sentinels(payload: object) -> set[str]:
    """Any embedded synthetic PII/secret sentinel string found anywhere in a
    JSON-shaped fragment (e.g. ``manifest.semantic_minima``), found via the
    same structural patterns ``scan_artifact`` itself checks for."""
    text = json.dumps(payload, ensure_ascii=False, default=str)
    values: set[str] = set()
    values.update(_EMAIL_RE.findall(text))
    values.update(match.group(0) for match in _BEARER_RE.finditer(text))
    values.update(_JWT_RE.findall(text))
    return values


def build_fixture_forbidden_values(manifest: JourneyManifest) -> frozenset[str]:
    """Derive every forbidden value from the real, loaded Task 1 manifest.

    Sources: the actual synthetic cell/paragraph content of every referenced
    input file, every literal string in the dialogue scenarios (minus
    citation IDs -- see ``_collect_dialogue_values``), every synthetic
    target/hash identifier in the governance outcomes and HITL plan
    instances, and any PII/secret-shaped string embedded in the semantic
    minima. Governance/plan branch names (``id``/``branch``/
    ``execution_class``/``expected_status``) are deliberately excluded: they
    are the same small enum ``ArtifactAllowlist.state_transitions`` is meant
    to legitimately carry (``approved``/``rejected``/``expired``/
    ``automatic``), so forbidding them would make every passing scan
    impossible. There is no hand-maintained literal list here.
    """
    values: set[str] = set()

    for entry in manifest.inputs:
        values.update(_extract_input_values(entry))

    for scenario in manifest.dialogue_scenarios:
        values.update(_collect_dialogue_values(scenario))

    for outcome in manifest.governance_outcomes:
        for key in ("required_plan_hash", "target_before_hash", "target_after_hash"):
            value = outcome.get(key)
            if isinstance(value, str) and value:
                values.add(value)

    for plan in manifest.plan_instances:
        for key in ("target_fixture_id", "target_before_hash", "target_after_hash"):
            value = plan.get(key)
            if isinstance(value, str) and value:
                values.add(value)

    values.update(_collect_pii_secret_sentinels(manifest.semantic_minima))

    return frozenset(values)


# ---------------------------------------------------------------------------
# Browser artifact redaction.
# ---------------------------------------------------------------------------


def _contains_forbidden(data: bytes, forbidden_values: frozenset[str]) -> bool:
    text = data.decode("utf-8", errors="replace")
    return any(value and value in text for value in forbidden_values)


def _infer_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return "screenshot"
    if suffix in {".zip", ".trace"}:
        return "trace"
    return "log"


def _infer_content_type(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".zip": "application/zip",
        ".log": "text/plain",
    }.get(suffix, "application/octet-stream")


def sanitize_browser_artifacts(
    paths: Sequence[Path], *, forbidden_values: frozenset[str]
) -> tuple[SanitizedBrowserArtifactRef, ...]:
    """Emit only redacted metadata for browser artifacts; omit unsafe ones.

    Raw Playwright traces, screenshots, and logs are always unsafe to
    upload as-is: this only ever returns a hash/size/kind reference, and
    only when the underlying file contains no forbidden value.
    """
    refs: list[SanitizedBrowserArtifactRef] = []
    for path in paths:
        try:
            data = path.read_bytes()
        except OSError:
            continue
        if _contains_forbidden(data, forbidden_values):
            continue
        refs.append(
            SanitizedBrowserArtifactRef(
                kind=_infer_kind(path),
                original_name_hash=sha256_text(path.name),
                size_bytes=len(data),
                content_type=_infer_content_type(path),
            )
        )
    return tuple(refs)


# ---------------------------------------------------------------------------
# Scanning.
# ---------------------------------------------------------------------------


def _has_sensitive_header_field(node: object) -> bool:
    if isinstance(node, Mapping):
        for key, value in node.items():
            if isinstance(key, str) and key.strip().lower() in _SENSITIVE_HEADER_KEYS:
                return True
            if _has_sensitive_header_field(value):
                return True
        return False
    if isinstance(node, (list, tuple)):
        return any(_has_sensitive_header_field(item) for item in node)
    return False


def scan_artifact(path: Path, *, forbidden_values: frozenset[str]) -> None:
    """Recursively scan one file's bytes/fields; raise on anything unsafe.

    Checks (in order): fixture-manifest-derived forbidden value echoes,
    PII email patterns, bearer-token/secret patterns, JWT-shaped values, and
    sensitive header field names in any parsed JSON structure. The raised
    message never embeds the raw matched value, the file path, or file
    contents -- only a small, fixed set of category tokens.
    """
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")

    reasons: set[str] = set()
    if any(value and value in text for value in forbidden_values):
        reasons.add("FORBIDDEN_VALUE_ECHOED")
    if _EMAIL_RE.search(text):
        reasons.add("PII_EMAIL_DETECTED")
    if _BEARER_RE.search(text):
        reasons.add("SECRET_BEARER_TOKEN_DETECTED")
    if _JWT_RE.search(text):
        reasons.add("JWT_LIKE_VALUE_DETECTED")

    try:
        payload = json.loads(text)
    except ValueError:
        payload = None
    if payload is not None and _has_sensitive_header_field(payload):
        reasons.add("SENSITIVE_HEADER_FIELD_DETECTED")

    if reasons:
        raise ArtifactSafetyError(",".join(sorted(reasons)))


# `run.json` (`orchestrator.STAGING_RELATIVE_PATH`) lives in the SAME
# staging directory the CI scan reads (Task 3's prepare phase writes it
# there for the browser to read) but is not `ArtifactAllowlist`-shaped --
# it is its own, already-allowlisted staging-manifest schema (`orchestrator.
# _reject_forbidden_manifest_content`). It is still scanned for forbidden
# values/PII/secrets like every other staged file; it is just exempt from
# the (different) `ArtifactAllowlist` schema check and never copied into the
# sanitized upload directory, which only ever carries per-journey CI
# evidence.
NON_ALLOWLIST_STAGING_FILENAMES = frozenset({"run.json"})


def _validate_allowlist_schema(payload: Mapping[str, object]) -> None:
    try:
        ArtifactAllowlist.model_validate(payload)
    except ValidationError as exc:
        raise ArtifactSafetyError("ARTIFACT_SCHEMA_VIOLATION") from exc


def _validate_staged_file_identity(path: Path, staging_dir: Path, payload: Mapping[str, object], *, run_id: str) -> None:
    """A file that individually passes ``ArtifactAllowlist`` (right schema,
    no forbidden values) can still be a real journey's evidence staged
    under the WRONG name -- e.g. `finance.json` actually carrying
    `credit`'s content, or a stale file from an entirely different run.
    Nothing previously cross-checked a staged file's own self-declared
    `run_id`/`journey_id` against the run being scanned or the filename
    `assemble_journey_evidence` is required to have written it under
    (`<staging>/<journey_id>.json`, the same convention
    `_missing_journey_evidence` already relies on)."""
    if payload.get("run_id") != run_id:
        raise ArtifactSafetyError("ARTIFACT_RUN_ID_MISMATCH")
    if path.parent == staging_dir and path.stem in JOURNEY_IDS and payload.get("journey_id") != path.stem:
        raise ArtifactSafetyError("ARTIFACT_JOURNEY_ID_MISMATCH")


def scan_and_materialize(
    staging_dir: Path,
    *,
    output_dir: Path,
    failure_summary_path: Path,
    manifest: JourneyManifest | None = None,
    forbidden_values: frozenset[str] | None = None,
    run_id: str,
) -> ScanResult:
    """Scan every file under ``staging_dir`` and materialize the outcome.

    On success: writes each staged file, re-validated and re-serialized
    through the closed ``ArtifactAllowlist`` schema, into a fresh
    ``output_dir``. On failure: writes only a fixed-schema
    ``ScannerFailureSummary`` to ``failure_summary_path`` and creates no
    output directory.

    Exactly one of ``manifest`` (single-journey forbidden values, the
    original per-journey call shape every existing test in this module
    uses) or ``forbidden_values`` (an already-unioned set, used by the CI
    ``scan`` subcommand below, which scans all three journeys' staged
    evidence in one pass) must be supplied.
    """
    if forbidden_values is None:
        if manifest is None:
            raise ValueError("scan_and_materialize requires either manifest or forbidden_values")
        forbidden_values = build_fixture_forbidden_values(manifest)
    staged_files = sorted(p for p in staging_dir.rglob("*") if p.is_file())

    counts = {"scanned": 0, "failed": 0}
    reason_codes: set[str] = set()

    for path in staged_files:
        counts["scanned"] += 1
        failed_this_file = False
        try:
            scan_artifact(path, forbidden_values=forbidden_values)
        except ArtifactSafetyError as exc:
            reason_codes.update(str(exc).split(","))
            failed_this_file = True
        if path.name not in NON_ALLOWLIST_STAGING_FILENAMES:
            try:
                payload = json.loads(path.read_text(encoding="utf-8", errors="replace"))
                _validate_allowlist_schema(payload)
                _validate_staged_file_identity(path, staging_dir, payload, run_id=run_id)
            except ArtifactSafetyError as exc:
                reason_codes.update(str(exc).split(","))
                failed_this_file = True
            except ValueError:
                reason_codes.add("ARTIFACT_SCHEMA_VIOLATION")
                failed_this_file = True
        if failed_this_file:
            counts["failed"] += 1

    if reason_codes:
        summary = ScannerFailureSummary(
            schema_version=1,
            run_id=run_id,
            status="scanner_failed",
            reason_codes=tuple(sorted(reason_codes)),
            category_counts={code: 1 for code in reason_codes},
            scanner_version=SCANNER_VERSION,
        )
        failure_summary_path.parent.mkdir(parents=True, exist_ok=True)
        failure_summary_path.write_text(json.dumps(summary.to_json(), indent=2, sort_keys=True))
        # Restores the documented contract ("on failure it removes/does not
        # create that [sanitized] directory") even when this is the SECOND
        # scan of the same staging directory in one run (the CI gate's own
        # final phase, then the workflow's separate `id: scan` step) and an
        # earlier, now-stale successful scan already created one.
        if output_dir.exists():
            shutil.rmtree(output_dir, ignore_errors=True)
        return ScanResult(
            status="failed",
            safe_evidence_dir=None,
            failure_summary_path=failure_summary_path,
            scan_safe=False,
            reason_codes=tuple(sorted(reason_codes)),
            counts=counts,
        )

    # Always a fresh, atomically-rebuilt directory -- never a merge with
    # whatever a previous call left behind. This tolerates the CI gate's own
    # two real invocations of this scan (its final phase, then the
    # workflow's separate `id: scan` step) over identical, unchanged staged
    # content without silently trusting stale output from the first call.
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=False)
    for path in staged_files:
        if path.name in NON_ALLOWLIST_STAGING_FILENAMES:
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        allowed = ArtifactAllowlist.model_validate(payload)
        target = output_dir / path.relative_to(staging_dir)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(allowed.model_dump_json(indent=2))

    return ScanResult(
        status="passed",
        safe_evidence_dir=output_dir,
        failure_summary_path=failure_summary_path,
        scan_safe=True,
        reason_codes=(),
        counts=counts,
    )


def scan_one_journey(argv: Sequence[str] | None = None) -> int:
    """The original single-journey CLI (kept for direct/manual use; the CI
    gate itself calls the ``scan`` subcommand in ``main`` below, which scans
    all three journeys' staged evidence in one pass)."""
    parser = argparse.ArgumentParser(description="Scan and materialize one journey's staged CI evidence.")
    parser.add_argument("staging_dir", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--failure-summary", type=Path, required=True)
    parser.add_argument("--journey-id", required=True, choices=JOURNEY_IDS)
    parser.add_argument(
        "--runtime-root",
        type=Path,
        default=_RUNTIME_DATA_DIR,
        help="Root containing <journey_id>/manifest.json (defaults to test_data/runtime).",
    )
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args(argv)

    manifest = load_journey_manifest(args.journey_id, args.runtime_root)
    result = scan_and_materialize(
        args.staging_dir,
        output_dir=args.output,
        failure_summary_path=args.failure_summary,
        manifest=manifest,
        run_id=args.run_id,
    )
    print(f"scan_safe={'true' if result.scan_safe else 'false'}")
    return 0 if result.scan_safe else 1


# ---------------------------------------------------------------------------
# All-journeys CI scan (the real gate's own entrypoint).
# ---------------------------------------------------------------------------


def build_all_journeys_forbidden_values(runtime_root: Path) -> frozenset[str]:
    """Union every journey's own forbidden values -- the CI staging
    directory holds all three journeys' evidence side by side, so no single
    journey's manifest is enough to scan it."""
    values: set[str] = set()
    for journey_id in JOURNEY_IDS:
        manifest = load_journey_manifest(journey_id, runtime_root)
        values.update(build_fixture_forbidden_values(manifest))
    return frozenset(values)


def assemble_journey_evidence(
    *,
    run_id: str,
    journey_id: str,
    preparation: Mapping[str, object],
    verification: Mapping[str, object],
) -> ArtifactAllowlist:
    """Project one journey's real, persisted prepare+verify records
    (``JourneyPreparation.to_dict()`` / ``JourneyVerification.to_dict()``,
    Task 3) into the closed ``ArtifactAllowlist`` schema for CI staging.

    ``verify_journey`` (``orchestrator.py``) already reconstructs the FULL
    per-journey totals into ``JourneyVerification`` -- ``logical_model_calls``,
    ``http_attempts``, ``retry_count``, and ``call_kinds`` there already cover
    all three logical calls (the preparation-phase ontology call, replayed
    from the staging run manifest's real recorded counters, plus the two
    browser-driven calls; see ``_ontology_call_from_manifest``/`` calls = (
    ontology_call,) + evidence.model_calls`` in that module), as do
    ``model_caller``/``model_origin``/``preflight_model_id``/
    ``requested_model_id``/``observed_model_ids``. Only ``preparation`` is
    consulted here, for the one field verification does not carry:
    ``fixture_manifest_sha256``. Adding preparation's own (already-included)
    counters on top would double the budget (previously a real bug fixed
    here: 4 logical calls instead of the mandated 3).

    Disclosed, honest gap: ``call_timestamps`` is always empty -- a per-call
    wall-clock timestamp exists on the low-level ``ModelResponse`` dataclass
    (``contracts.py``), but is never propagated up through
    ``ModelCallRecord``/``JourneyVerification``, so there is genuinely
    nothing to read at this level; this is not fabricated as a placeholder.
    ``backend_trace_ids`` IS populated for real, from ``verification``'s own
    ``correlation_ids`` (``"{run_id}:{journey_id}:{call_kind}:{index}"`` per
    logical call, already computed by ``verify_journey``) -- a genuine,
    non-fabricated backend-assigned trace identifier per call, not an empty
    placeholder. ``semantic_outcomes`` is not a separately-measured field
    either; it is safe to mark every logical call "passed" here because this
    function is only ever called after `prepare_journey`'s semantic-minimum
    validation and `verify_journey`'s structural checks have both already
    raised on any failure -- by the time a preparation/verification record
    reaches here, every one of its calls passed.
    """
    logical_model_calls = int(verification["logical_model_calls"])
    observed_model_ids = tuple(verification["observed_model_ids"])
    return ArtifactAllowlist(
        schema_version=1,
        run_id=run_id,
        journey_id=journey_id,
        fixture_manifest_sha256=str(preparation["fixture_manifest_sha256"]),
        model_requested=str(verification["requested_model_id"]),
        model_observed=str(observed_model_ids[-1]) if observed_model_ids else "",
        model_origin=str(verification["model_origin"]),
        model_caller=str(verification["model_caller"]),
        model_config_version_id=str(verification["model_config_version_id"]),
        preflight_model_id=str(verification["preflight_model_id"]),
        logical_model_calls=logical_model_calls,
        http_attempts=int(verification["http_attempts"]),
        retry_count=int(verification["retry_count"]),
        call_timestamps=(),
        call_kinds=tuple(verification["call_kinds"]),
        semantic_outcomes=("passed",) * logical_model_calls,
        state_transitions=tuple(branch["branch"] for branch in verification["plan_branches"]),
        citation_ids=tuple(verification["citation_ids"]),
        tool_trace_ids=tuple(verification["tool_trace_ids"]),
        audit_event_ids=tuple(verification["audit_event_ids"]),
        backend_trace_ids=tuple(verification.get("correlation_ids") or ()),
    )


_DETERMINISTIC_REPORT_REASON = "DETERMINISTIC_REPORT_NOT_CLEAN"


def _check_deterministic_report(path: Path | None) -> str | None:
    """Return a reason code if ``path`` (``run_registered_cases.py``'s own
    report) does not show a clean, zero-model-call deterministic run;
    ``None`` if the report is clean or was not supplied."""
    if path is None:
        return None
    if not path.exists():
        return _DETERMINISTIC_REPORT_REASON
    try:
        report = json.loads(path.read_text(encoding="utf-8"))
    except ValueError:
        return _DETERMINISTIC_REPORT_REASON
    if report.get("missing") or report.get("duplicates") or report.get("model_calls", 0) != 0:
        return _DETERMINISTIC_REPORT_REASON
    for result in report.get("results", []):
        if result.get("status") != "passed" or result.get("skipped"):
            return _DETERMINISTIC_REPORT_REASON
    return None


_JOURNEY_EVIDENCE_MISSING_REASON = "JOURNEY_EVIDENCE_MISSING"


def _missing_journey_evidence(staging_dir: Path) -> tuple[str, ...]:
    """Journeys with no ``<staging>/<journey_id>.json`` (`assemble_journey_
    evidence`'s own output, written by ``run.py`` after a passing verify).
    Without this check, a staging directory containing only Task 3's
    `run.json` (e.g. because an earlier phase failed before any journey's
    evidence file was written) would scan zero files, hit no forbidden
    value, and vacuously report `scan_safe=true` -- exactly backwards for a
    gate whose `id: scan` step runs with `if: always()` specifically to
    still produce a safe result after an earlier failure."""
    return tuple(
        journey_id for journey_id in JOURNEY_IDS
        if not (staging_dir / f"{journey_id}.json").exists()
    )


def _write_scanner_failure(failure_summary_path: Path, *, run_id: str, reason_codes: tuple[str, ...]) -> None:
    summary = ScannerFailureSummary(
        schema_version=1,
        run_id=run_id,
        status="scanner_failed",
        reason_codes=reason_codes,
        category_counts={code: 1 for code in reason_codes},
        scanner_version=SCANNER_VERSION,
    )
    failure_summary_path.parent.mkdir(parents=True, exist_ok=True)
    failure_summary_path.write_text(json.dumps(summary.to_json(), indent=2, sort_keys=True))


_RUN_ID_UNRESOLVED = "unresolved-run-id"


def _resolve_run_id(explicit: str | None, staging_dir: Path) -> str:
    """Never raises: the CI gate's own `id: scan` step runs with `if:
    always()` specifically so it can still emit `scan_safe=false` and a
    failure summary after an earlier phase failed -- including a phase 3
    (prepare) failure so early that `<staging>/run.json` was never written.
    Raising here would make `actions/upload-artifact@v4`'s failure-summary
    step error on a file that also never got written, instead of cleanly
    uploading one. Falling through to `_missing_journey_evidence` (which
    will always also be true when `run.json` itself is missing) is what
    actually produces the failure summary."""
    if explicit:
        return explicit
    run_manifest = staging_dir / "run.json"
    if run_manifest.exists():
        try:
            document = json.loads(run_manifest.read_text(encoding="utf-8"))
        except ValueError:
            return _RUN_ID_UNRESOLVED
        run_id = document.get("run_id")
        if run_id:
            return str(run_id)
    return _RUN_ID_UNRESOLVED


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="python -m evals.business_journeys.artifacts")
    subparsers = parser.add_subparsers(dest="command", required=True)

    scan = subparsers.add_parser(
        "scan", help="Scan all three journeys' staged CI evidence and materialize the sanitized/failure output.",
    )
    scan.add_argument("--staging", type=Path, required=True)
    scan.add_argument("--sanitized-output", type=Path, required=True)
    scan.add_argument("--failure-summary", type=Path, required=True)
    scan.add_argument(
        "--manifest", type=Path, default=None,
        help="test_data/runtime/manifest.json; its parent directory is the runtime root "
        "the three per-journey manifests load from (defaults to test_data/runtime).",
    )
    scan.add_argument(
        "--deterministic-report", type=Path, default=None,
        help="run_registered_cases.py's own report; if given, the scan only passes when it "
        "shows zero missing/duplicate/failed/skipped cases and zero model calls.",
    )
    scan.add_argument("--run-id", default=None, help="defaults to <staging>/run.json's run_id")
    return parser


def _run_scan(args: argparse.Namespace) -> int:
    runtime_root = args.manifest.parent if args.manifest else _RUNTIME_DATA_DIR
    run_id = _resolve_run_id(args.run_id, args.staging)

    # Checked before scanning a single byte: a staging directory missing a
    # journey's evidence file (e.g. because an earlier phase failed first)
    # must never vacuously pass just because it also contains no forbidden
    # value.
    missing = _missing_journey_evidence(args.staging)
    if missing:
        _write_scanner_failure(args.failure_summary, run_id=run_id, reason_codes=(_JOURNEY_EVIDENCE_MISSING_REASON,))
        print("scan_safe=false")
        return 1

    deterministic_reason = _check_deterministic_report(args.deterministic_report)

    forbidden_values = build_all_journeys_forbidden_values(runtime_root)
    result = scan_and_materialize(
        args.staging,
        output_dir=args.sanitized_output,
        failure_summary_path=args.failure_summary,
        forbidden_values=forbidden_values,
        run_id=run_id,
    )

    if deterministic_reason is not None and result.scan_safe:
        # The evidence itself was safe, but the deterministic registry run
        # this gate depends on was not clean -- fail closed the same way a
        # forbidden-value hit would, with the same fixed-schema summary, and
        # remove the sanitized directory `scan_and_materialize` just created
        # so a not-actually-safe run leaves no sanitized evidence behind.
        if result.safe_evidence_dir is not None:
            shutil.rmtree(result.safe_evidence_dir, ignore_errors=True)
        _write_scanner_failure(args.failure_summary, run_id=run_id, reason_codes=(deterministic_reason,))
        print("scan_safe=false")
        return 1

    print(f"scan_safe={'true' if result.scan_safe else 'false'}")
    return 0 if result.scan_safe else 1


def main(argv: Sequence[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)
    if args.command == "scan":
        return _run_scan(args)
    parser.error(f"unknown command {args.command!r}")
    return 2


if __name__ == "__main__":
    sys.exit(main())
